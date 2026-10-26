# -*- coding: utf-8 -*-
"""Summary of every validated sample: RESULTS/_validation/<modality>_summary.xlsx
(+ .json).

Sheets: Cases (one row per sample, every metric of every alignment),
Descriptive (per model: n, mean, std, median, Q1, Q3), Per phantom (same, per
model and phantom, primary alignment), Tests (models compared on the samples
they all have: Wilcoxon signed-rank for 2 models, Friedman + pairwise
Wilcoxon with Holm correction for 3 or more; needs scipy).

Only the newest run of each (sample, model) is kept (date folder).
"""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List

import numpy as np

#: Metrics compared between models (primary alignment).
KEY_METRICS = ("accuracy.rmse_mm", "accuracy.mean_signed_mm", "accuracy.pct_lt_1mm",
               "accuracy.pct_lt_2mm", "completeness.median_abs_mm", "completeness.pct_lt_1mm",
               "completeness.pct_lt_2mm", "hausdorff_p95_mm", "fiducials.norm_rmse_mm",
               "fiducials.normal_rmse_mm", "fiducials.in_plane_rmse_mm")


def _date(name: str):
    try:
        return datetime.strptime(name, "%d-%m-%Y")
    except (TypeError, ValueError):
        return datetime.min


def collect(results_root: Path, modality: str = "vision") -> List[dict]:
    newest: Dict[tuple, dict] = {}
    for p in sorted(Path(results_root).glob(f"*/*/validation/{modality}/metrics.json")):
        try:
            r = json.loads(p.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            print(f"[validation] skipped {p}: {exc}")
            continue
        r["date"] = p.parents[3].name
        key = (r.get("sample"), r.get("model"))
        if key not in newest or _date(r["date"]) > _date(newest[key]["date"]):
            newest[key] = r
    return list(newest.values())


def _flat(prefix: str, d: Any, out: dict) -> None:
    if isinstance(d, dict):
        for k, v in d.items():
            _flat(f"{prefix}{k}.", v, out)
    elif isinstance(d, (int, float)) and not isinstance(d, bool):
        out[prefix[:-1]] = float(d)


def row_of(r: dict) -> dict:
    row = {k: r.get(k) for k in ("sample", "phantom", "pose", "rep", "model", "date",
                                 "primary_alignment")}
    row["kabsch.corner_rms_mm"] = r["alignment"]["kabsch"]["corner_rms_mm"]
    for name, m in r["metrics"].items():
        flat: dict = {}
        _flat("", {k: v for k, v in m.items() if k != "height"}, flat)
        for k, v in flat.items():
            row[f"{name}.{k}"] = v
        h = m.get("height", {})
        if "slope_mm_per_mm" in h:
            row[f"{name}.height.slope_mm_per_mm"] = h["slope_mm_per_mm"]
    for k in ("backend_s", "total_s"):
        if (r.get("timings") or {}).get(k) is not None:
            row[f"time.{k}"] = float(r["timings"][k])
    v = r.get("views") or {}
    for k in ("n_views", "n_views_seeing_phantom", "azimuth_coverage_deg"):
        if k in v:
            row[f"views.{k}"] = float(v[k])
    for k in ("distance_mm", "elevation_deg"):
        for s, x in (v.get(k) or {}).items():
            row[f"views.{k}.{s}"] = float(x)
    return row


def _desc(x: np.ndarray) -> dict:
    x = x[np.isfinite(x)]
    if x.size == 0:
        return {"n": 0}
    q1, med, q3 = np.percentile(x, [25, 50, 75])
    return {"n": int(x.size), "mean": float(x.mean()),
            "std": float(x.std(ddof=1)) if x.size > 1 else 0.0,
            "median": float(med), "q1": float(q1), "q3": float(q3)}


def _holm(ps: List[float]) -> List[float]:
    order = np.argsort(ps)
    adj, prev = [0.0] * len(ps), 0.0
    for rank, i in enumerate(order):
        prev = max(prev, min(1.0, (len(ps) - rank) * ps[i]))
        adj[i] = prev
    return adj


def _block(r: dict):
    """Same physical case across models: (phantom, pose, rep), else the sample name."""
    key = (r.get("phantom"), r.get("pose"), r.get("rep"))
    return key if all(key) else (r.get("sample"),)


def tests(rows: List[dict], primary_of: Dict[str, str]) -> List[dict]:
    models = sorted({r["model"] for r in rows})
    if len(models) < 2:
        return [{"note": f"one model only ({', '.join(models)}): no comparison"}]
    try:
        from scipy import stats
    except ImportError:
        return [{"note": "scipy not installed: no statistical tests"}]
    by_model = {m: {_block(r): r for r in rows if r["model"] == m} for m in models}
    blocks = sorted(set.intersection(*(set(v) for v in by_model.values())), key=str)

    def wilcoxon(a, b):
        return float(stats.wilcoxon(a, b).pvalue) if np.any(a != b) else 1.0

    out = []
    for metric in KEY_METRICS:
        data = {m: np.array([by_model[m][b].get(f"{primary_of[m]}.{metric}", np.nan)
                             for b in blocks], dtype=float) for m in models}
        ok = np.all([np.isfinite(v) for v in data.values()], axis=0) if blocks else np.zeros(0, bool)
        n = int(ok.sum())
        if n < 3:
            out.append({"metric": metric, "note": f"only {n} samples common to all models"})
            continue
        if len(models) == 2:
            a, b = (data[m][ok] for m in models)
            out.append({"metric": metric, "test": "Wilcoxon", "models": " vs ".join(models),
                        "n": n, "p": wilcoxon(a, b)})
            continue
        cols = [data[m][ok] for m in models]
        if all(np.array_equal(cols[0], c) for c in cols[1:]):
            out.append({"metric": metric, "note": "identical values for every model"})
            continue
        p = float(stats.friedmanchisquare(*cols).pvalue)
        out.append({"metric": metric, "test": "Friedman", "models": ", ".join(models), "n": n, "p": p})
        pairs = [(a, b) for i, a in enumerate(models) for b in models[i + 1:]]
        raw = [wilcoxon(data[a][ok], data[b][ok]) for a, b in pairs]
        for (a, b), pr, pa in zip(pairs, raw, _holm(raw)):
            out.append({"metric": metric, "test": "Wilcoxon (Holm)", "models": f"{a} vs {b}",
                        "n": n, "p": pr, "p_holm": pa})
    return out


def _sheet(wb, title: str, rows: List[dict]) -> None:
    from openpyxl.styles import Font, PatternFill

    ws = wb.create_sheet(title)
    cols: List[str] = []
    for r in rows:
        cols += [k for k in r if k not in cols]
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", start_color="305496")
    for r in rows:
        ws.append([round(v, 5) if isinstance(v, float) else v for v in (r.get(c) for c in cols)])
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(10, min(40, len(c) + 2))
    ws.freeze_panes = "B2"


def summarize(cfg, modality: str = "vision") -> Path:
    from openpyxl import Workbook

    results = collect(cfg.results_root, modality)
    if not results:
        raise SystemExit(f"[validation] no {modality} validation under {cfg.results_root}"
                         f"/*/*/validation/{modality}/: run `spectra validate` first")
    rows = [row_of(r) for r in results]
    primary_of: Dict[str, str] = {}
    for r in results:
        primary_of.setdefault(r["model"], r.get("primary_alignment", "kabsch"))
    numeric = sorted({k for r in rows for k, v in r.items() if isinstance(v, float)})

    desc, per_phantom = [], []
    for m in sorted({r["model"] for r in rows}):
        rm = [r for r in rows if r["model"] == m]
        for k in numeric:
            desc.append({"model": m, "metric": k,
                         **_desc(np.array([r.get(k, np.nan) for r in rm], dtype=float))})
        for ph in sorted({r["phantom"] for r in rm}, key=str):
            rp = [r for r in rm if r["phantom"] == ph]
            for k in KEY_METRICS:
                key = f"{primary_of[m]}.{k}"
                per_phantom.append({"model": m, "phantom": ph, "metric": key,
                                    **_desc(np.array([r.get(key, np.nan) for r in rp], dtype=float))})
    test_rows = tests(rows, primary_of)

    out_dir = Path(cfg.results_root) / "_validation"
    out_dir.mkdir(parents=True, exist_ok=True)
    xlsx = out_dir / f"{modality}_summary.xlsx"
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Cases", rows)
    _sheet(wb, "Descriptive", desc)
    _sheet(wb, "Per phantom", per_phantom)
    _sheet(wb, "Tests", test_rows)
    wb.save(xlsx)
    (out_dir / f"{modality}_summary.json").write_text(json.dumps(
        {"cases": rows, "descriptive": desc, "per_phantom": per_phantom, "tests": test_rows},
        indent=2), encoding="utf-8")
    print(f"[validation] {len(rows)} samples, {len({r['model'] for r in rows})} model(s) -> {xlsx}")
    return xlsx
