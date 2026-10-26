# -*- coding: utf-8 -*-
"""Ground truth in the CAD frame (mm, Z up): the cassette's ArUco markers,
from the marker table, the phantom STL and its fiducial holes (the bersagli,
found in the STL itself: ``find_holes``).

Marker table (Tabella_marker_cassetta.xlsx), sheet ``Marker``: one row per
marker with the centre of its black square, the height of its visible
surface and the direction its printed top points to (0 = +Y, 90 = -X,
180 = -Y, 270 = +X; counter-clockwise seen from above). Sheet ``Info``:
nominal / measured edge of the black square. The 4 corners are returned in
the ArUco order (top-left, top-right, bottom-right, bottom-left of the
printed marker), the order vision writes in ``aruco_markers_3d.json``.
"""

from __future__ import annotations

import math
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Sequence

import numpy as np

from .geometry import to_h

#: Columns of the `Marker` sheet (matched by prefix, case-insensitive).
REQUIRED_COLS = ("marker_id", "x_centro_mm", "y_centro_mm", "z_superficie_mm", "rotazione_deg")
EDGE_COL = "lato_misurato_mm"


@dataclass
class Fixture:
    corners_mm: Dict[int, np.ndarray]   # marker id -> (4, 3), ArUco order, CAD frame
    edge_mm: float                      # edge used for the markers without their own
    edge_source: str                    # "measured" | "nominal"
    path: Path
    warnings: List[str] = field(default_factory=list)


def _num(value) -> Optional[float]:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, str):
        value = value.strip().replace(",", ".")
    return float(value)


def marker_corners(cx: float, cy: float, z: float, rot_deg: float, edge: float) -> np.ndarray:
    """Corners TL, TR, BR, BL of a marker centred in (cx, cy) at height z whose
    printed top points along +Y rotated by rot_deg (counter-clockwise from above)."""
    h = edge / 2.0
    base = np.array([[-h, h], [h, h], [h, -h], [-h, -h]])      # top towards +Y
    a = math.radians(rot_deg)
    R = np.array([[math.cos(a), -math.sin(a)], [math.sin(a), math.cos(a)]])
    xy = base @ R.T + np.array([cx, cy])
    return np.c_[xy, np.full(4, float(z))]


def read_fixture(path: Path) -> Fixture:
    from openpyxl import load_workbook

    path = Path(path)
    if not path.is_file():
        raise SystemExit(f"[validation] marker table not found: {path}")
    wb = load_workbook(path, data_only=True)
    info = {}
    if "Info" in wb.sheetnames:
        for row in wb["Info"].iter_rows(values_only=True):
            if row and row[0] is not None and len(row) > 1:
                info[str(row[0]).strip().lower()] = row[1]

    def info_num(prefix):
        for k, v in info.items():
            if k.startswith(prefix):
                return _num(v)
        return None

    measured, nominal = info_num("lato misurato"), info_num("lato nominale")
    edge, source = (measured, "measured") if measured else (nominal, "nominal")
    warnings = [] if measured else ["marker edge not measured (Info sheet): using the nominal one"]

    if "Marker" not in wb.sheetnames:
        raise SystemExit(f"[validation] {path.name}: sheet 'Marker' missing")
    rows = list(wb["Marker"].iter_rows(values_only=True))
    header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
    col = {}
    for name in REQUIRED_COLS + (EDGE_COL, "note"):
        hits = [i for i, h in enumerate(header) if h.startswith(name)]
        if hits:
            col[name] = hits[0]
    missing = [c for c in REQUIRED_COLS if c not in col]
    if missing:
        raise SystemExit(f"[validation] {path.name}: columns {missing} missing in sheet 'Marker'")

    corners: Dict[int, np.ndarray] = {}
    for r_i, row in enumerate(rows[1:], start=2):
        if row is None or all(v is None for v in row):
            continue
        note = str(row[col["note"]] or "") if "note" in col else ""
        if "esempio" in note.lower() or row[col["marker_id"]] is None:
            continue
        mid = int(_num(row[col["marker_id"]]))
        vals = [_num(row[col[c]]) for c in REQUIRED_COLS[1:]]
        if any(v is None for v in vals):
            raise SystemExit(f"[validation] {path.name}, row {r_i} (marker {mid}): empty values")
        e = _num(row[col[EDGE_COL]]) if EDGE_COL in col else None
        e = e or edge
        if not e:
            raise SystemExit(f"[validation] {path.name}: marker edge missing (Info sheet or "
                             f"lato_misurato_mm of marker {mid})")
        if mid in corners:
            raise SystemExit(f"[validation] {path.name}: marker {mid} listed twice")
        corners[mid] = marker_corners(*vals, e)
    if not corners:
        raise SystemExit(f"[validation] {path.name}: no markers in sheet 'Marker'")
    return Fixture(corners, float(edge or 0.0), source, path, warnings)


def cad_frame(corners_mm: Dict[int, np.ndarray], *, marker_ids: Sequence[int],
              origin_ids: Optional[Sequence[int]], x_ids: Optional[Sequence[int]]):
    """CAD -> output-frame transform built with the SAME rule as vision
    (``spectra.vision.align.fit_marker_plane``), Z = CAD +Z. Returns (T 4x4, info)."""
    from ..vision.align import MarkerTriangulation, fit_marker_plane

    ms = {}
    for i in marker_ids:
        c = np.asarray(corners_mm[i], dtype=np.float64)
        ms[int(i)] = MarkerTriangulation(
            marker_id=int(i), corners_3d=c, num_views=0, reproj_rmse_px=0.0,
            edge_lengths_3d=np.linalg.norm(np.roll(c, -1, axis=0) - c, axis=1),
            center_3d=c.mean(axis=0))
    frame, info = fit_marker_plane(ms, reference_up=np.array([0.0, 0.0, 1.0]),
                                   frame_marker_ids=origin_ids, x_axis_marker_ids=x_ids)
    R = np.asarray(frame.basis, dtype=np.float64).T
    return to_h(R, -R @ np.asarray(frame.origin, dtype=np.float64)), info


def load_mesh(path: Path):
    """(vertices, faces) of a mesh file, as float64 / int64."""
    import trimesh

    path = Path(path)
    if not path.is_file():
        raise SystemExit(f"[validation] mesh not found: {path}")
    m = trimesh.load(path, force="mesh", process=False)
    return np.asarray(m.vertices, dtype=np.float64), np.asarray(m.faces, dtype=np.int64)


# ------------------------------------------------------- fiducial holes ---
@dataclass
class Holes:
    """Blind round holes of the phantom (the bersagli), CAD frame, mm."""
    mouth: np.ndarray      # (N, 3) centre of the hole mouth, on the phantom surface
    axis: np.ndarray       # (N, 3) unit, out of the hole (outward surface normal there)
    diameter: np.ndarray   # (N,)
    depth: np.ndarray      # (N,)

    def __len__(self) -> int:
        return len(self.mouth)


def find_holes(vertices: np.ndarray, faces: np.ndarray, max_diameter_mm: float = 3.0,
               tol: float = 0.05) -> Holes:
    """Blind round holes on the flat faces of a mesh: groups of small planar
    regions (walls + floor) whose vertices all lie on a cylinder around the
    floor normal, above the floor. Mouth = floor centre + depth * axis. Sorted
    by Y then X (CAD); the ids in the outputs are 1-based positions here."""
    import trimesh
    from scipy.sparse import coo_matrix
    from scipy.sparse.csgraph import connected_components

    m = trimesh.Trimesh(np.asarray(vertices, dtype=np.float64), np.asarray(faces, dtype=np.int64),
                        process=True)                       # merges the duplicated STL vertices
    n_f = len(m.faces)
    region = np.arange(n_f) + len(m.facets)                 # planar regions: facets + lone faces
    for k, f in enumerate(m.facets):
        region[f] = k
    area = np.bincount(region, weights=m.area_faces)
    small = area[region] < math.pi * max_diameter_mm ** 2 / 4.0
    adj = m.face_adjacency
    ok = small[adj[:, 0]] & small[adj[:, 1]]
    graph = coo_matrix((np.ones(int(ok.sum())), (adj[ok, 0], adj[ok, 1])), shape=(n_f, n_f))
    _, comp = connected_components(graph, directed=False)

    found = []
    for c in np.unique(comp[small]):
        fs = np.flatnonzero((comp == c) & small)
        regs = np.unique(region[fs])
        if len(regs) < 4:                                    # a floor and >= 3 walls
            continue
        floor = fs[region[fs] == regs[np.argmax(area[regs])]]
        w = m.area_faces[floor]
        axis = (m.face_normals[floor] * w[:, None]).sum(axis=0)
        axis /= np.linalg.norm(axis)
        centre = (m.triangles_center[floor] * w[:, None]).sum(axis=0) / w.sum()
        walls = np.setdiff1d(fs, floor)
        v = m.vertices[np.unique(m.faces[walls])] - centre   # (a floor may have a centre vertex)
        h = v @ axis
        radial = np.linalg.norm(v - h[:, None] * axis, axis=1)
        r = float(np.median(radial))
        depth = float(h.max())
        if (not 0 < 2 * r <= max_diameter_mm
                or np.abs(radial - r).max() > tol * r          # not a cylinder around the axis
                or not 0.85 < w.sum() / (math.pi * r * r) < 1.0 + tol   # floor not a disc
                or h.min() < -tol * r or depth < tol * r):     # not a blind hole above its floor
            continue
        found.append((centre + depth * axis, axis, 2.0 * r, depth))
    found.sort(key=lambda x: (round(float(x[0][1]), 2), float(x[0][0])))
    if not found:
        return Holes(np.zeros((0, 3)), np.zeros((0, 3)), np.zeros(0), np.zeros(0))
    mouth, ax, d, dep = (np.array(x) for x in zip(*found))
    return Holes(mouth, ax, d, dep)
